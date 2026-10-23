"""Allowlisted selections; financial dates select identities, never clip lifetime."""

import io
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, localcontext
from functools import wraps
from zoneinfo import ZoneInfo

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q, Sum

from . import clock
from .models import Animal, Audit, Event, Inventory, Property


def precise(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        with localcontext() as context:
            context.prec = 60
            return fn(*args, **kwargs)

    return wrapped


DATASETS = [
    "animals",
    "inventory",
    "feed",
    "treatments",
    "restrictions",
    "arrivals",
    "exits",
    "audit",
    "lpa2",
    "lpa3C",
    "lpa3D",
    "lpa5A",
    "lpa5B",
    "lpa6",
]
FIELDS = [
    "animal",
    "identity",
    "status",
    "breed",
    "sex",
    "age_weeks",
    "estimated_age",
    "mob",
    "purchase",
    "sale",
    "feed",
    "chemical",
    "procedure",
    "vet",
    "cost",
    "revenue",
    "profit",
    "financial_complete",
    "tags",
    "restriction_end",
    "effective_at",
    "submitted_at",
    "event",
    "kind",
    "name",
    "balance",
    "unit",
    "value",
    "details",
    "actor",
    "action",
    "count",
]
FILTERS = {
    "breed",
    "sex",
    "status",
    "animal",
    "mob",
    "min_age",
    "max_age",
    "present_from",
    "present_to",
    "activity_from",
    "activity_to",
    "as_of",
}


def validate_definition(definition):
    if not isinstance(definition, dict) or set(definition) - {
        "dataset",
        "columns",
        "filters",
        "sort",
        "group",
        "attribute_time",
        "filter_mode",
    }:
        raise ValidationError(
            "Unsupported report settings. Executable expressions and SQL are not accepted."
        )
    if definition.get("dataset", "animals") not in DATASETS:
        raise ValidationError("Unsupported report dataset.")
    if (
        set(definition.get("columns", [])) - set(FIELDS)
        or definition.get("sort", "animal").lstrip("-") not in FIELDS
        or definition.get("group", "") not in ["", "mob", "breed", "sex", "status", "kind", "name"]
    ):
        raise ValidationError("Choose allowed report columns, grouping and sort fields.")
    filters = definition.get("filters", {})
    if not isinstance(filters, dict) or set(filters) - FILTERS:
        raise ValidationError("Unsupported report filter.")
    for key in FILTERS - {"breed", "sex", "status", "animal", "mob"}:
        if filters.get(key) not in [None, ""]:
            try:
                int(filters[key])
            except (TypeError, ValueError):
                raise ValidationError("Report ages/timestamps must be integers.")
    if definition.get("attribute_time", "current") not in [
        "current",
        "as_of",
    ] or definition.get("filter_mode", "all") not in ["all", "any"]:
        raise ValidationError("Invalid attribute time or filter combination.")
    if definition.get("attribute_time") == "as_of" and not filters.get("as_of"):
        raise ValidationError("An as-of timestamp is required for historical attributes.")
    return definition


@precise
def animal_rows(actor, definition):
    from .models import AttributeChange

    f = definition.get("filters", {})
    at = int(f.get("as_of") or clock.now())
    rows = []
    for a in Animal.objects.filter(property=actor.property, purged=False).prefetch_related(
        "identifiers", "allocations", "memberships__mob", "restrictions"
    ):
        breed, sex, status = a.breed, a.sex, a.status
        if definition.get("attribute_time") == "as_of":
            for field in ["breed", "sex"]:
                later = (
                    AttributeChange.objects.filter(
                        animal=a,
                        field=field,
                        event__effective_at__gt=at,
                        event__reversed_by__isnull=True,
                    )
                    .order_by("event__effective_at", "event__effective_order")
                    .first()
                )
                if later:
                    if field == "breed":
                        breed = later.before
                    else:
                        sex = later.before
            p = a.presence.filter(start__lte=at).order_by("-start").first()
            status = "active" if p and (p.end is None or p.end > at) else "absent"
        age = None
        if a.dob:
            age = (
                datetime.fromtimestamp(at, ZoneInfo(actor.property.timezone)).date() - a.dob
            ).days // 7
        elif a.estimated_weeks is not None and a.age_reference is not None:
            age = (
                a.estimated_weeks
                + (
                    datetime.fromtimestamp(at, ZoneInfo(actor.property.timezone)).date()
                    - datetime.fromtimestamp(
                        a.age_reference, ZoneInfo(actor.property.timezone)
                    ).date()
                ).days
                // 7
            )
        mob = (
            a.memberships.filter(superseded=False, start__lte=at)
            .filter(Q(end__isnull=True) | Q(end__gt=at))
            .select_related("mob")
            .first()
        )
        checks = []
        if f.get("breed"):
            checks.append(breed.casefold() == str(f["breed"]).casefold())
        if f.get("sex"):
            checks.append(sex.casefold() == str(f["sex"]).casefold())
        if f.get("status"):
            checks.append(status == f["status"])
        if f.get("animal"):
            checks.append(
                str(a.id) == str(f["animal"])
                or a.identifiers.filter(value=str(f["animal"]).upper()).exists()
            )
        if f.get("mob"):
            checks.append(a.memberships.filter(mob_id=f["mob"], superseded=False).exists())
        if f.get("min_age"):
            checks.append(age is not None and age >= int(f["min_age"]))
        if f.get("max_age"):
            checks.append(age is not None and age <= int(f["max_age"]))
        if f.get("present_from") or f.get("present_to"):
            lo = int(f.get("present_from") or 0)
            hi = int(f.get("present_to") or clock.now())
            checks.append(
                a.presence.filter(start__lte=hi)
                .filter(Q(end__isnull=True) | Q(end__gte=lo))
                .exists()
            )
        if f.get("activity_from") or f.get("activity_to"):
            q = a.participant_set.filter(
                event__effective_at__gte=int(f.get("activity_from") or 0),
                event__effective_at__lte=int(f.get("activity_to") or clock.now()),
            )
            checks.append(q.exists())
        if checks and not (any(checks) if definition.get("filter_mode") == "any" else all(checks)):
            continue
        sums = {
            r["category"]: r["total"]
            for r in a.allocations.values("category").annotate(total=Sum("amount"))
        }
        revenue = sums.get("sale", Decimal(0))
        cost = -sum(v for k, v in sums.items() if k != "sale")
        complete = not a.allocations.filter(
            unknown=True, resolved_by__isnull=True, event__reversed_by__isnull=True
        ).exists()
        restrictions = (
            a.restrictions.filter(event__reversed_by__isnull=True, superseded_by__isnull=True)
            .order_by("-ends_at")
            .first()
        )
        row = {
            "animal": a.label,
            "identity": str(a.id),
            "status": status,
            "breed": breed or "Unknown",
            "sex": sex or "Unknown",
            "age_weeks": age,
            "estimated_age": not bool(a.dob),
            "mob": mob.mob.name if mob else "Unassigned / exited",
            "cost": cost,
            "revenue": revenue,
            "profit": revenue - cost if complete else None,
            "financial_complete": complete,
            "tags": "; ".join(
                f"{t.kind}:{t.value} ({'current' if t.current else 'retired'})"
                for t in a.identifiers.all()
            ),
            "restriction_end": restrictions.ends_at if restrictions else None,
        }
        for cat in ["purchase", "sale", "feed", "chemical", "procedure", "vet"]:
            row[cat] = sums.get(cat, Decimal(0)) if cat == "sale" else -sums.get(cat, Decimal(0))
        rows.append(row)
    return rows


@precise
@transaction.atomic
def build(actor, definition):
    validate_definition(definition)
    Property.objects.select_for_update().get(pk=actor.property_id)
    ds = definition.get("dataset", "animals")
    if ds in LPA_FIELDS:
        rows = lpa_records(actor, ds)
        definition = {
            **definition,
            "columns": definition.get("columns") or LPA_FIELDS[ds],
        }
    elif ds == "animals":
        rows = animal_rows(actor, definition)
    elif ds == "inventory":
        rows = [
            {
                "name": i.name,
                "kind": i.kind,
                "packages_remaining": i.packages_remaining,
                "package_size": i.package_size,
                "package_price": i.package_price,
                "cost_per_stock_unit": i.cost_per_stock_unit,
                "dose_unit": i.dose_unit,
                "cost_per_dose_unit": i.cost_per_dose_unit if i.kind == "chemical" else None,
                "batch": i.batch,
                "balance": i.balance,
                "unit": i.base_unit,
                "value": i.remaining_value,
                "details": {
                    "retired": i.retired,
                    "batch": i.batch,
                    "date_kind": i.date_kind,
                    "date_value": str(i.date_value or ""),
                    "package_size": str(i.package_size),
                    "package_price": str(i.package_price),
                },
            }
            for i in Inventory.objects.filter(property=actor.property)
        ]
    elif ds == "restrictions":
        from .models import Restriction

        rows = [
            {
                "animal": r.animal.label,
                "effective_at": r.starts_at,
                "restriction_end": r.ends_at,
                "details": {"combined_days": r.days, "timezone": r.timezone},
                "event": str(r.event_id),
            }
            for r in Restriction.objects.filter(
                animal__property=actor.property,
                event__reversed_by__isnull=True,
                superseded_by__isnull=True,
            )
        ]
    elif ds == "audit":
        rows = [
            {
                "actor": r.actor.username,
                "action": r.action,
                "submitted_at": r.recorded_at,
                "event": str(r.event_id or ""),
                "details": r.details,
            }
            for r in Audit.objects.filter(property=actor.property).select_related("actor")
        ]
    else:
        kinds = {
            "feed": ["feed"],
            "lpa3C": ["feed"],
            "lpa3D": ["receipt"],
            "lpa6": ["receipt"],
            "treatments": ["treat", "procedure", "vet", "observation"],
            "lpa2": ["treat"],
            "arrivals": ["register"],
            "lpa5A": ["register"],
            "exits": ["sale", "death"],
            "lpa5B": ["sale"],
        }
        rows = []
        for e in Event.objects.filter(
            property=actor.property,
            kind__in=kinds[ds]
            + (
                ["reverse", "allocation_correction"]
                if ds in ["feed", "treatments", "exits"]
                else []
            ),
        ).select_related("actor"):
            source_event = e
            if e.kind in ["reverse", "allocation_correction"]:
                original = Event.objects.filter(
                    property=actor.property, pk=e.details.get("original")
                ).first()
                if not original or original.kind not in kinds[ds]:
                    continue
                source_event = original
            amounts = list(e.allocations.values("category", "amount", "unknown", "resolved_by"))
            cost = -sum((v["amount"] for v in amounts if v["category"] != "sale"), Decimal(0))
            revenue = sum((v["amount"] for v in amounts if v["category"] == "sale"), Decimal(0))
            if (
                ds == "lpa3D"
                and e.details.get("kind") != "feed"
                or ds == "lpa6"
                and e.details.get("kind") != "chemical"
            ):
                continue
            rows.append(
                {
                    "event": str(e.id),
                    "kind": e.kind,
                    "mob": MobName(actor, source_event.details.get("mob")),
                    "cost": cost,
                    "revenue": revenue,
                    "profit": revenue - cost,
                    "effective_at": e.effective_at,
                    "submitted_at": e.submitted_at,
                    "actor": e.actor.username,
                    "count": e.participants.count(),
                    "details": {
                        **e.details,
                        "reversed_by": str(e.reversed_by_id or ""),
                        "record_status": "Reversed" if e.reversed_by_id else "Recorded",
                    },
                }
            )
        f = definition.get("filters", {})
        if f.get("activity_from"):
            rows = [r for r in rows if r["effective_at"] >= int(f["activity_from"])]
        if f.get("activity_to"):
            rows = [r for r in rows if r["effective_at"] <= int(f["activity_to"])]
    if ds != "animals":
        filters = definition.get("filters", {})
        if filters.get("activity_from"):
            rows = [
                r
                for r in rows
                if int(r.get("effective_at", r.get("submitted_at", 0)))
                >= int(filters["activity_from"])
            ]
        if filters.get("activity_to"):
            rows = [
                r
                for r in rows
                if int(r.get("effective_at", r.get("submitted_at", 0)))
                <= int(filters["activity_to"])
            ]
        animal_filters = {
            k: v for k, v in filters.items() if k not in ["activity_from", "activity_to"]
        }
        if set(animal_filters) - {"as_of"} and ds not in ["inventory", "audit", "lpa3D", "lpa6"]:
            selected = {
                r["identity"] for r in animal_rows(actor, {**definition, "filters": animal_filters})
            }
            event_ids = {
                str(v)
                for v in Event.objects.filter(
                    property=actor.property, participants__animal_id__in=selected
                ).values_list("id", flat=True)
            }
            rows = [
                r
                for r in rows
                if str(r.get("event", "")) in event_ids
                or str(r.get("animal", ""))
                in {
                    r["animal"]
                    for r in animal_rows(actor, {**definition, "filters": animal_filters})
                }
            ]
    sort = definition.get("sort", "animal")
    group = definition.get("group")

    def sort_value(row):
        value = row.get(sort.lstrip("-"))
        return (
            str(row.get(group, "")) if group else "",
            value is None,
            value if isinstance(value, (Decimal, int)) else str(value or ""),
        )

    rows.sort(key=sort_value, reverse=sort.startswith("-"))
    totals = {
        k: sum((r[k] for r in rows if isinstance(r.get(k), Decimal)), Decimal(0))
        for k in ["cost", "revenue", "profit"]
    }
    groups = []
    if group:
        from itertools import groupby

        for label, items in groupby(rows, key=lambda row: str(row.get(group, ""))):
            items = list(items)
            groups.append(
                {
                    "name": label,
                    "count": len(items),
                    "cost": sum(
                        (
                            r.get("cost", Decimal(0))
                            for r in items
                            if isinstance(r.get("cost"), Decimal)
                        ),
                        Decimal(0),
                    ),
                    "revenue": sum(
                        (
                            r.get("revenue", Decimal(0))
                            for r in items
                            if isinstance(r.get("revenue"), Decimal)
                        ),
                        Decimal(0),
                    ),
                    "profit": sum(
                        (
                            r.get("profit", Decimal(0))
                            for r in items
                            if isinstance(r.get("profit"), Decimal)
                        ),
                        Decimal(0),
                    ),
                }
            )
    return {
        "rows": rows,
        "groups": groups,
        "totals": totals,
        "definition": definition,
        "selection_description": selection_description(actor, definition),
        "generated_at": clock.now(),
        "display_timezone": actor.display_timezone,
        "display_format": actor.date_format,
        "note": "AUD. Date/attribute filters select animals; their finances always cover the full lifetime. Unknown purchase prices make P/L incomplete. Active animals show costs/revenue to date, not a final sale outcome. Export dates use your display timezone. Property timezone "
        + actor.property.timezone
        + ". Amounts are displayed to cents with reconciliation rounding; Excel also includes full-precision values.",
    }


def text(value):
    import json

    return (
        json.dumps(value, default=str, ensure_ascii=False)
        if isinstance(value, (dict, list))
        else ""
        if value is None
        else str(value)
    )


@precise
def xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "StockLOG"
    ws.append(["StockLOG report"])
    ws.append([report["note"]])
    ws.append(["Selection", report.get("selection_description", text(report["definition"]))])
    columns = report["definition"].get("columns") or (
        list(report["rows"][0]) if report["rows"] else ["animal", "cost", "revenue", "profit"]
    )
    ws.append(columns)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for index, row in enumerate(report["rows"]):
        ws.append([display_value(row.get(c), c, report, index) for c in columns])
        for cell in ws[ws.max_row]:
            # Explicit string cell type, not a formula, including leading =,+,-,@ text.
            cell.data_type = "s"
    ws.append(["Totals (full precision)", text(report["totals"])])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = (
        f"A4:{__import__('openpyxl').utils.get_column_letter(len(columns))}{ws.max_row - 1}"
    )
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(
            60, max(16, len(str(column[3].value or "")) + 2)
        )
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    if report.get("groups"):
        summary = wb.create_sheet("Groups")
        summary.append(["Group", "Records", "Cost AUD", "Revenue AUD", "Profit AUD"])
        for group in report["groups"]:
            summary.append(
                [
                    group["name"],
                    group["count"],
                    text(group["cost"]),
                    text(group["revenue"]),
                    text(group["profit"]),
                ]
            )
    raw = wb.create_sheet("Full precision")
    raw.append(["Stored precision in AUD/quantities; no financial lifetime clipping."])
    raw.append(columns)
    for row in report["rows"]:
        raw.append([text(row.get(c)) for c in columns])
        for cell in raw[raw.max_row]:
            cell.data_type = "s"
    raw.freeze_panes = "A3"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@precise
def pdf(report):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    story = [
        Paragraph(
            "StockLOG — " + report["definition"].get("dataset", "animals"),
            styles["Title"],
        ),
        Paragraph(escape(report["note"]), styles["Normal"]),
        Spacer(1, 8),
        Paragraph(
            "Selection: " + escape(report.get("selection_description", text(report["definition"]))),
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    default = (
        ["animal", "status", "cost", "revenue", "profit", "financial_complete"]
        if report["definition"].get("dataset", "animals") == "animals"
        else ["event", "kind", "effective_at", "cost", "revenue", "actor", "details"]
    )
    if report["definition"].get("dataset") == "inventory":
        default = ["name", "kind", "packages_remaining", "balance", "unit", "value", "batch"]
    elif report["definition"].get("dataset") == "restrictions":
        default = ["animal", "effective_at", "restriction_end", "details"]
    elif report["definition"].get("dataset") == "audit":
        default = ["actor", "action", "submitted_at", "details"]
    columns = report["definition"].get("columns") or default
    for offset in range(0, len(columns), 7):
        subset = columns[offset : offset + 7]
        if offset and columns[0] not in subset:
            subset = [columns[0]] + subset
        if offset:
            story.append(PageBreak())
        table = [
            [Paragraph(escape(c.replace("_", " ").title()), styles["BodyText"]) for c in subset]
        ]
        for index, row in enumerate(report["rows"]):
            table.append(
                [
                    Paragraph(
                        escape(display_value(row.get(c), c, report, index)).replace("\n", "<br/>"),
                        styles["BodyText"],
                    )
                    for c in subset
                ]
            )
        t = Table(
            table,
            colWidths=[(landscape(A4)[0] - 48) / len(subset)] * len(subset),
            repeatRows=1,
            hAlign="LEFT",
            splitInRow=1,
        )
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d7e7dc")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(t)
    if report.get("groups"):
        story.append(Spacer(1, 12))
        for group in report["groups"]:
            story.append(
                Paragraph(
                    escape(
                        "Group "
                        + group["name"]
                        + ": "
                        + text({k: v for k, v in group.items() if k != "name"})
                    ),
                    styles["Normal"],
                )
            )
    story += [
        Spacer(1, 12),
        Paragraph(
            "Totals (AUD, full precision): " + escape(text(report["totals"])),
            styles["Normal"],
        ),
    ]
    doc.build(story)
    return out.getvalue()


# Each included LPA section has explicit export columns rather than an opaque JSON cell.
LPA_FIELDS = {
    "lpa2": [
        "treatment_date",
        "livestock_description",
        "location",
        "identifiers_or_mob",
        "livestock_count",
        "product_name",
        "batch",
        "expiry_date",
        "date_kind",
        "actual_dose_per_head",
        "dose_unit",
        "administration_method",
        "rate_description",
        "approximate_live_weight",
        "administered_by",
        "administered_by_contact",
        "submitting_user",
        "combined_withholding_esi_days",
        "safe_for_slaughter_at",
        "adverse_reactions",
        "broken_needle_status",
        "broken_needle_details",
        "equipment_cleaned_calibrated",
        "equipment_description",
        "equipment_checked_on",
        "equipment_checked_by_name",
        "equipment_checked_by_contact",
        "expired_override",
        "override_reason",
    ],
    "lpa3C": [
        "feedstuff_description",
        "storage_location",
        "mob",
        "feeding_start",
        "feeding_finish",
        "responsible_name",
        "submitting_user",
        "quantity",
        "quantity_mode",
        "consumed",
        "cost",
    ],
    "lpa3D": [
        "date_received",
        "feedstuff",
        "amount_packages",
        "quantity_per_package",
        "stock_unit",
        "received_quantity",
        "origin",
        "supplier",
        "declaration_type",
        "declaration_reference",
        "declaration_details",
        "certificate_available",
        "product_tested",
        "residue_test_reference",
        "storage_location",
        "signed_by",
        "submitting_user",
    ],
    "lpa5A": [
        "arrival_date",
        "nvd",
        "nlis_upload_reference",
        "stock_count",
        "destination_paddock",
        "breed",
        "sex",
        "vendor_name",
        "vendor_address",
        "source_pic",
        "notes",
        "identifiers",
    ],
    "lpa5B": [
        "exit_date",
        "nvd",
        "nlis_upload_reference",
        "stock_count",
        "origin_paddock",
        "breed",
        "sex",
        "destination_name",
        "destination_pic",
        "identifiers",
        "restriction_acceptance",
        "revenue",
    ],
    "lpa6": [
        "inventory_category",
        "date_received",
        "product_name",
        "quantity",
        "stock_unit",
        "date_value",
        "date_kind",
        "batch",
        "updated_by",
        "submitting_user",
    ],
}
for _fields in LPA_FIELDS.values():
    _fields.append("record_status")
    for _field in _fields:
        if _field not in FIELDS:
            FIELDS.append(_field)


def lpa_records(actor, dataset):
    kinds = {
        "lpa2": ["treat"],
        "lpa3C": ["feed"],
        "lpa3D": ["receipt"],
        "lpa5A": ["register"],
        "lpa5B": ["sale"],
        "lpa6": ["receipt"],
    }
    records = []
    for e in Event.objects.filter(property=actor.property, kind__in=kinds[dataset]).select_related(
        "actor"
    ):
        d = e.details
        if (
            dataset == "lpa3D"
            and d.get("kind") != "feed"
            or dataset == "lpa6"
            and d.get("kind") != "chemical"
        ):
            continue
        labels = "; ".join(
            p.snapshot.get("label", "[retained identity]") for p in e.participants.all()
        )
        base = {
            "event": str(e.id),
            "effective_at": e.effective_at,
            "submitted_at": e.submitted_at,
            "kind": e.kind,
            "submitting_user": e.actor.username,
            "reversed_by": str(e.reversed_by_id or ""),
            "record_status": "Reversed" if e.reversed_by_id else "Recorded",
        }
        if dataset == "lpa2":
            end = RestrictionEnd(e)
            r = {
                "treatment_date": e.effective_at,
                "livestock_description": d.get("livestock_description", ""),
                "location": d.get("location_text", ""),
                "identifiers_or_mob": labels,
                "livestock_count": d.get("count", 0),
                "product_name": d.get("product_name", ""),
                "batch": d.get("batch", ""),
                "expiry_date": d.get("date_value", "")
                if d.get("date_kind") == "expiry"
                else "Not supplied as expiry",
                "date_kind": d.get("date_kind", ""),
                "actual_dose_per_head": d.get("dose", ""),
                "dose_unit": d.get("dose_unit", ""),
                "approximate_live_weight": "Not collected — owner scope override",
                "combined_withholding_esi_days": d.get("combined_restriction_days", 0),
                "safe_for_slaughter_at": end,
                "administered_by": d.get("administered_by", e.actor.username),
            }
            for k in [
                "administered_by_contact",
                "administration_method",
                "rate_description",
                "adverse_reactions",
                "broken_needle_status",
                "broken_needle_details",
                "equipment_cleaned_calibrated",
                "equipment_description",
                "equipment_checked_on",
                "equipment_checked_by_name",
                "equipment_checked_by_contact",
                "expired_override",
                "override_reason",
            ]:
                r[k] = d.get(k, "Unknown" if k.endswith("_status") else "")
        elif dataset == "lpa3C":
            item = Inventory.objects.get(pk=d["item"], property=actor.property)
            r = {
                "feedstuff_description": item.name,
                "storage_location": d.get("storage", item.storage),
                "mob": MobName(actor, d.get("mob")),
                "feeding_start": e.effective_at,
                "feeding_finish": e.effective_at,
                "responsible_name": d.get("responsible_name", e.actor.username),
                "quantity": d.get("quantity"),
                "quantity_mode": d.get("mode"),
                "consumed": d.get("consumed"),
                "cost": d.get("cost"),
            }
        elif dataset == "lpa3D":
            r = {
                "date_received": e.effective_at,
                "feedstuff": d.get("name"),
                "amount_packages": d.get("packages"),
                "quantity_per_package": d.get("package_size"),
                "stock_unit": d.get("stock_unit"),
                "received_quantity": d.get("received"),
                "origin": d.get("origin_text", ""),
                "supplier": d.get("supplier_name", ""),
                "certificate_available": d.get("residue_certificate_available", "Unknown"),
                "product_tested": d.get("product_tested", "Unknown"),
                "storage_location": d.get("storage", ""),
                "signed_by": d.get("signed_by_name", ""),
            }
            for k in [
                "declaration_type",
                "declaration_reference",
                "declaration_details",
                "residue_test_reference",
            ]:
                r[k] = d.get(k, "")
        elif dataset == "lpa5A":
            for row in d.get("rows") or [d]:
                r = {
                    "arrival_date": row.get("effective_at", e.effective_at),
                    "nvd": row.get("nvd", ""),
                    "nlis_upload_reference": row.get("nlis_upload_reference", ""),
                    "stock_count": 1,
                    "destination_paddock": row.get("destination_paddock_text", ""),
                    "breed": row.get("breed", "Unknown"),
                    "sex": row.get("sex", "Unknown"),
                    "vendor_name": row.get("vendor_name", ""),
                    "vendor_address": row.get("vendor_address", ""),
                    "source_pic": row.get("source_pic", ""),
                    "notes": row.get("notes", ""),
                    "identifiers": "; ".join(str(row.get(k, "")) for k in ["nlis", "rfid"]),
                }
                records.append({**base, **r})
            continue
        elif dataset == "lpa5B":
            snapshots = list(e.participants.values_list("snapshot", flat=True))
            r = {
                "exit_date": e.effective_at,
                "nvd": d.get("nvd", ""),
                "nlis_upload_reference": d.get("nlis_upload_reference", ""),
                "stock_count": d.get("count", len(snapshots)),
                "origin_paddock": d.get("origin_paddock_text", ""),
                "breed": "; ".join(sorted({v.get("breed", "Unknown") for v in snapshots})),
                "sex": "; ".join(sorted({v.get("sex", "Unknown") for v in snapshots})),
                "destination_name": d.get("destination_name", ""),
                "destination_pic": d.get("destination_pic", ""),
                "identifiers": labels,
                "restriction_acceptance": d.get("restriction_acceptance", []),
                "revenue": d.get("revenue"),
            }
        else:
            r = {
                "inventory_category": d.get("category", "veterinary"),
                "date_received": e.effective_at,
                "product_name": d.get("name"),
                "quantity": d.get("received"),
                "stock_unit": Inventory.objects.get(pk=d["item"]).base_unit,
                "date_value": d.get("date_value", ""),
                "date_kind": d.get("date_kind", ""),
                "batch": d.get("batch", ""),
                "updated_by": d.get("updated_by_name") or e.actor.username,
            }
        records.append({**base, **r})
    return records


def RestrictionEnd(event):
    from .models import Restriction

    return Restriction.objects.filter(event=event, superseded_by__isnull=True).aggregate(
        end=__import__("django").db.models.Max("ends_at")
    )["end"]


def MobName(actor, pk):
    from .models import Mob

    return Mob.objects.get(pk=pk, property=actor.property).name if pk else ""


MONEY_FIELDS = {
    "purchase",
    "sale",
    "feed",
    "chemical",
    "procedure",
    "vet",
    "cost",
    "revenue",
    "profit",
    "value",
}
TIME_FIELDS = {
    "effective_at",
    "submitted_at",
    "restriction_end",
    "treatment_date",
    "safe_for_slaughter_at",
    "feeding_start",
    "feeding_finish",
    "date_received",
    "arrival_date",
    "exit_date",
}


@precise
def display_value(value, column, report, index):
    if column in TIME_FIELDS and value not in [None, ""]:
        try:
            return datetime.fromtimestamp(
                int(value), ZoneInfo(report.get("display_timezone", "Australia/Melbourne"))
            ).strftime(report.get("display_format", "%d/%m/%Y %H:%M"))
        except (ValueError, TypeError, OverflowError):
            return text(value)
    if column in MONEY_FIELDS and isinstance(value, Decimal):
        from decimal import ROUND_FLOOR

        cache = report.setdefault("_display_money", {})
        if column not in cache:
            cent = Decimal("0.01")
            values = [
                (i, r.get(column))
                for i, r in enumerate(report["rows"])
                if isinstance(r.get(column), Decimal)
            ]
            floors = {i: v.quantize(cent, rounding=ROUND_FLOOR) for i, v in values}
            goal = sum((v for i, v in values), Decimal(0)).quantize(cent, rounding=ROUND_HALF_UP)
            extra = int((goal - sum(floors.values())) / cent)
            ordered = sorted(
                values,
                key=lambda pair: (
                    -(pair[1] - floors[pair[0]]),
                    str(report["rows"][pair[0]].get("identity", pair[0])),
                ),
            )
            for i, v in ordered[:extra]:
                floors[i] += cent
            cache[column] = floors
        return format(cache[column][index], ".2f")
    return text(value)


def selection_description(actor, definition):
    pieces = [definition.get("dataset", "animals").replace("_", " ").title()]
    for key, value in definition.get("filters", {}).items():
        if key == "mob":
            try:
                value = MobName(actor, value)
            except Exception:
                value = "Unavailable mob"
        elif key in ["as_of", "present_from", "present_to", "activity_from", "activity_to"]:
            value = display_value(
                value,
                "effective_at",
                {
                    "rows": [],
                    "display_timezone": actor.display_timezone,
                    "display_format": actor.date_format,
                },
                0,
            )
        pieces.append(key.replace("_", " ").title() + ": " + str(value))
    pieces.append("Attributes: " + definition.get("attribute_time", "current").replace("_", " "))
    pieces.append("Filters: " + definition.get("filter_mode", "all"))
    return "; ".join(pieces)


for _field in [
    "packages_remaining",
    "package_size",
    "package_price",
    "cost_per_stock_unit",
    "dose_unit",
    "cost_per_dose_unit",
    "batch",
    "record_status",
]:
    if _field not in FIELDS:
        FIELDS.append(_field)
