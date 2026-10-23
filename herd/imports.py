import csv
import io
import zipfile
from datetime import date, datetime

from django.core.exceptions import ValidationError

from .clock import local_timestamp
from .services import preview

MAX_ROWS = 2000
COLUMNS = [
    "nlis",
    "rfid",
    "effective_at",
    "source_pic",
    "nvd",
    "purchase_cost",
    "breed",
    "sex",
    "estimated_weeks",
    "dob",
    "mob",
    "vendor_name",
    "vendor_address",
    "nlis_upload_reference",
    "destination_paddock_text",
]


def read_upload(upload):
    if upload.size > 8 * 1024 * 1024:
        raise ValidationError("Import limit is 8 MiB.")
    raw = upload.read()
    if upload.name.lower().endswith(".csv"):
        rows = []
        for row in csv.reader(io.StringIO(raw.decode("utf-8-sig"))):
            if len(rows) > MAX_ROWS:
                raise ValidationError("Import exceeds 2000 rows.")
            if len(row) > 80:
                raise ValidationError("Import exceeds 80 columns.")
            rows.append(row)
    elif upload.name.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if (
                len(z.infolist()) > 500
                or sum(x.file_size for x in z.infolist()) > 32 * 1024 * 1024
                or any("vbaProject" in x.filename for x in z.infolist())
            ):
                raise ValidationError("Spreadsheet archive is unsafe or too large.")
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        rows = []
        for row in wb.active.iter_rows():
            if len(rows) > MAX_ROWS:
                raise ValidationError("Import exceeds 2000 rows.")
            if len(row) > 80:
                raise ValidationError("Import exceeds 80 columns.")
            values = []
            for cell in row:
                if cell.data_type == "f":
                    raise ValidationError(
                        "Import formulas are not accepted. Export literal values first."
                    )
                value = cell.value
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                values.append("" if value is None else str(value))
            if any(values):
                rows.append(values)
        wb.close()
    else:
        raise ValidationError("Use CSV or XLSX structured data, not document attachments.")
    if len(rows) < 2 or len(rows) > MAX_ROWS + 1 or len(rows[0]) > 80:
        raise ValidationError("Import must have headers and 1–2000 rows, up to 80 columns.")
    return rows[0], rows[1:]


def mapped_preview(actor, headers, rows, mapping, shared):
    output = []
    errors = []
    seen = set()
    for index, source in enumerate(rows):
        d = {
            k: source[int(v)]
            for k, v in mapping.items()
            if v not in [None, ""] and int(v) < len(source)
        }
        d = {
            **{k: v for k, v in shared.items() if v not in [None, ""]},
            **{k: v for k, v in d.items() if v not in [None, ""]},
        }
        d["entry_source"] = "purchase"
        err = []
        try:
            if d.get("effective_at") and not str(d["effective_at"]).isdigit():
                d["effective_at"] = local_timestamp(d["effective_at"], actor.property.timezone)
            if not d.get("effective_at"):
                raise ValidationError("Arrival date is required.")
            for key in ["nlis", "rfid"]:
                val = d.get(key)
                if val:
                    pair = (key, str(val).strip().upper())
                    if pair in seen:
                        raise ValidationError("Duplicate identifier inside this import.")
                    seen.add(pair)
            preview(actor, "register", d)
        except (ValidationError, ValueError, TypeError) as e:
            err = getattr(e, "messages", [str(e)])
        missing = [
            k for k in ["purchase_cost", "breed", "sex", "mob", "estimated_weeks"] if not d.get(k)
        ]
        output.append({"number": index + 1, "data": d, "missing": missing, "errors": err})
        errors.append(err)
    return output, errors
