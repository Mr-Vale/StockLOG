import uuid
from decimal import Decimal as D

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Sum

from herd import clock, reports
from herd import services as s
from herd.models import (
    Allocation,
    Animal,
    Event,
    Identifier,
    Inventory,
    ReportTemplate,
)

from .test_foundations import arrival, run

T = 1600000000


def setup_mob(user, count=5):
    mob = run(user, "mob", {"name": "Test mob"})["mob"]
    for i in range(count):
        arrival(user, tag=f"EDGE{i:012d}", mob=mob, rfid=f"{i:015d}")
    return mob


def stock(user, kind="feed", **kw):
    d = {
        "kind": kind,
        "name": "Hay" if kind == "feed" else "Treatment",
        "packages": "10",
        "package_size": "330" if kind == "feed" else "100",
        "package_price": "65" if kind == "feed" else "50",
        "stock_unit": "kg" if kind == "feed" else "mL",
        **kw,
    }
    return run(user, "receipt", d)["item"]


@pytest.mark.django_db
def test_tag_lineage_and_reversal_reserves_forever(farm):
    _, a, b, c = farm
    with clock.at(T):
        result = arrival(b)
        animal = Animal.objects.get(pk=result["animals"][0])
        change = run(
            b,
            "replace_tag",
            {
                "animal": "A" * 16,
                "nlis": "N" * 16,
                "rfid": "000000000000099",
                "reason": "Lost tag",
            },
        )
        assert s.find_animal(c, "A" * 16).id == animal.id == s.find_animal(c, "N" * 16).id
        run(b, "reverse", {"original": change["event"], "reason": "Wrong replacement"})
        assert animal.identifiers.get(value="A" * 16).current
        assert not animal.identifiers.get(value="N" * 16).current
        with pytest.raises(ValidationError):
            arrival(b, tag="N" * 16)
        with pytest.raises(ValidationError):
            run(
                b,
                "replace_tag",
                {
                    "animal": "A" * 16,
                    "rfid": "１２３４５６７８９０１２３４５",
                    "reason": "Test",
                },
            )


@pytest.mark.django_db
def test_feed_cross_negative_refill_and_fractional_no_double_charge(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b)
        item = stock(b)
        r = run(b, "feed", {"mob": mob, "item": item, "mode": "bales", "quantity": "12"})
        assert Inventory.objects.get(pk=item).balance == D(-660) and D(r["cost"]) == 780
        with pytest.raises(ValidationError):
            run(c, "feed", {"mob": mob, "item": item, "mode": "bales", "quantity": "1"})
        assert stock(c) == item
        assert Inventory.objects.get(pk=item).balance == 2640
        assert -Allocation.objects.filter(category="feed").aggregate(v=Sum("amount"))["v"] == 780
        half = run(b, "feed", {"mob": mob, "item": item, "mode": "bales", "quantity": "0.5"})
        assert D(half["consumed"]) == 165 and D(half["cost"]) == D("32.5")
        run(b, "stocktake", {"item": item, "actual": 0, "reason": "Measured zero"})
        with pytest.raises(ValidationError):
            run(b, "feed", {"mob": mob, "item": item, "quantity": "1"})
        before = Allocation.objects.count()
        with pytest.raises(ValidationError):
            run(b, "stocktake", {"item": item, "actual": 100})
        assert Allocation.objects.count() == before


@pytest.mark.django_db
def test_undo_mistake_and_24_hour_permission(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b)
        item = stock(b)
        mistake = run(b, "feed", {"mob": mob, "item": item, "quantity": 3, "mode": "bales"})
    with clock.at(T + 3600):
        with pytest.raises(PermissionDenied):
            run(c, "reverse", {"original": mistake["event"], "reason": "Not own"})
        run(
            b,
            "reverse",
            {"original": mistake["event"], "reason": "Three was a mistake"},
        )
        corrected = run(b, "feed", {"mob": mob, "item": item, "quantity": 2, "mode": "bales"})
    with clock.at(T + 3600 + 86400):
        with pytest.raises(PermissionDenied):
            run(b, "reverse", {"original": corrected["event"], "reason": "Too late"})
        assert Inventory.objects.get(pk=item).balance == 2640
        assert (
            -sum(Allocation.objects.filter(category="feed").values_list("amount", flat=True)) == 130
        )
        run(a, "reverse", {"original": corrected["event"], "reason": "Admin correction"})


@pytest.mark.django_db
def test_chemical_shortage_units_expiry_and_dom(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b)
        item = stock(b, kind="chemical", packages=1, package_size=8, package_price=4)
        count = Event.objects.count()
        with pytest.raises(ValidationError):
            run(b, "treat", {"mob": mob, "item": item, "dose": 2})
        assert Inventory.objects.get(pk=item).balance == 8 and Event.objects.count() == count
        large = stock(
            b,
            kind="chemical",
            name="20 litre product",
            package_size=20,
            package_price=250,
            stock_unit="L",
            packages=1,
        )
        inv = Inventory.objects.get(pk=large)
        assert inv.unit_cost == D("0.0125") and inv.balance == 20000
        expired = stock(
            b,
            kind="chemical",
            name="Expired",
            date_kind="expiry",
            date_value="2020-01-01",
            packages=1,
        )
        with pytest.raises(ValidationError):
            run(b, "treat", {"mob": mob, "item": expired, "dose": 1})
        run(
            b,
            "treat",
            {
                "mob": mob,
                "item": expired,
                "dose": 1,
                "expired_override": True,
                "override_reason": "Actual historical use explicitly recorded",
            },
        )
        dom = stock(
            b,
            kind="chemical",
            name="DOM only",
            date_kind="DOM",
            date_value="2019-01-01",
            packages=1,
        )
        run(b, "treat", {"mob": mob, "item": dom, "dose": 1})


@pytest.mark.django_db
def test_historical_membership_and_admin_correction(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 2)
        item = stock(b)
        other = run(b, "mob", {"name": "Other"})["mob"]
    with clock.at(T + 200):
        run(b, "move", {"animal": "EDGE000000000000", "destination": other})
    with clock.at(T + 300):
        before = run(
            b,
            "feed",
            {
                "mob": mob,
                "item": item,
                "quantity": 1,
                "mode": "bales",
                "effective_at": T + 100,
            },
        )
        after = run(
            b,
            "feed",
            {
                "mob": mob,
                "item": item,
                "quantity": 1,
                "mode": "bales",
                "effective_at": T + 250,
            },
        )
        assert before["count"] == 2 and after["count"] == 1
        with pytest.raises(PermissionDenied):
            run(
                b,
                "move",
                {
                    "animal": "EDGE000000000001",
                    "destination": other,
                    "effective_at": T + 150,
                    "administration": True,
                },
            )
        reviewed = s.preview(
            a,
            "move",
            {
                "animal": "EDGE000000000000",
                "destination": other,
                "effective_at": T + 50,
                "administration": True,
            },
        )
        assert reviewed["effects"]["corrected_events"]
        result = s.execute(
            a,
            uuid.uuid4().hex,
            "move",
            {
                "animal": "EDGE000000000000",
                "destination": other,
                "effective_at": T + 50,
                "administration": True,
            },
            reviewed["version"],
        )
        assert result["corrected_events"]
        assert Event.objects.filter(kind="feed").count() == 2


@pytest.mark.django_db
def test_sale_death_restriction_repurchase_and_locks(farm):
    _, a, b, c = farm
    with clock.at(T):
        setup_mob(b, 2)
        chemical = stock(b, kind="chemical", restriction_days=28)
        run(b, "treat", {"animal": "EDGE000000000000", "item": chemical, "dose": 1})
    with clock.at(T + 100):
        data = {
            "animal": "EDGE000000000000",
            "destination_pic": "DST",
            "nvd": "NVD",
            "price": "2000",
        }
        with pytest.raises(ValidationError):
            run(b, "sale", data)
        sale = run(b, "sale", {**data, "accept_restriction": True})
        with pytest.raises(ValidationError):
            run(b, "treat", {"animal": "EDGE000000000000", "item": chemical, "dose": 1})
        with pytest.raises(PermissionDenied):
            run(b, "reverse", {"original": sale["event"], "reason": "Not admin"})
        death = run(b, "death", {"animal": "EDGE000000000001", "reason": "Synthetic"})
        with pytest.raises(PermissionDenied):
            run(b, "reverse", {"original": death["event"], "reason": "Not admin"})
        run(
            a,
            "reverse",
            {"original": death["event"], "reason": "Admin verified correction"},
        )
    with clock.at(T + 200):
        repurchase = arrival(
            b, tag="EDGE000000000000", entry_source="repurchase", purchase_cost="150"
        )
        row = reports.build(b, {"filters": {"animal": "EDGE000000000000"}})["rows"][0]
        assert row["purchase"] == 250 and row["revenue"] == 2000 and row["status"] == "active"
        assert (
            s.find_animal(b, "EDGE000000000000").id
            == Animal.objects.get(pk=repurchase["animals"][0]).id
        )


@pytest.mark.django_db
def test_unknown_cost_birth_missing_data_and_reporting(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b, purchase_cost="", sex="", breed="")
        row = reports.build(b, {})["rows"][0]
        assert not row["financial_complete"] and row["profit"] is None and row["age_weeks"] is None
        with pytest.raises(ValidationError):
            run(b, "register", {"nlis": "B" * 16, "entry_source": "birth"})
        run(
            b,
            "register",
            {"nlis": "B" * 16, "entry_source": "birth", "mother_nlis": "A" * 16},
        )
        born = reports.build(b, {"filters": {"animal": "B" * 16}})["rows"][0]
        assert born["financial_complete"] and born["purchase"] == 0
        mob = run(b, "mob", {"name": "Empty"})["mob"]
        item = stock(b)
        with pytest.raises(ValidationError):
            run(b, "feed", {"mob": mob, "item": item, "quantity": 1})


@pytest.mark.django_db
def test_report_lifetime_templates_injection_and_filters(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 2)
        item = stock(b)
    with clock.at(T + 100):
        run(b, "feed", {"mob": mob, "item": item, "quantity": 1, "mode": "bales"})
    with clock.at(T + 200):
        report = reports.build(
            b,
            {
                "filters": {"present_from": T, "present_to": T + 1, "breed": "Unknown"},
                "filter_mode": "any",
            },
        )
        assert report["totals"]["cost"] == 265
        create = run(
            b,
            "template",
            {
                "name": "Shared",
                "definition": {"dataset": "animals", "filters": {"mob": str(mob)}},
            },
        )
        run(
            c,
            "template",
            {
                "template": create["template"],
                "name": "Changed",
                "definition": {"dataset": "inventory"},
            },
        )
        run(
            b,
            "template",
            {
                "template": create["template"],
                "name": "Deleted",
                "definition": {},
                "delete": True,
            },
        )
        assert not ReportTemplate.objects.exists()
        with pytest.raises(ValidationError):
            reports.build(b, {"filters": {"sql": "DROP TABLE"}})
        import io

        from openpyxl import load_workbook

        malicious = {
            "rows": [{"animal": '=HYPERLINK("bad")'}],
            "totals": {},
            "definition": {},
            "note": "Synthetic",
            "generated_at": T,
        }
        ws = load_workbook(io.BytesIO(reports.xlsx(malicious))).active
        assert ws["A5"].data_type == "s" and ws["A5"].value.startswith("=")


@pytest.mark.django_db
def test_receipt_matching_retirement_and_adjustment_preserve_history(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 1)
        item = stock(b)
        assert stock(b) == item
        assert stock(b, package_price=66) != item
        run(b, "feed", {"mob": mob, "item": item, "quantity": 1, "mode": "bales"})
        cost = Allocation.objects.filter(category="feed").aggregate(v=Sum("amount"))["v"]
        run(b, "stocktake", {"item": item, "actual": 300, "reason": "Measured"})
        assert Allocation.objects.filter(category="feed").aggregate(v=Sum("amount"))["v"] == cost
        run(b, "retire", {"item": item})
        assert Inventory.objects.get(pk=item).retired
        with pytest.raises(ValidationError):
            run(b, "feed", {"mob": mob, "item": item, "quantity": 1})
        assert reports.build(b, {"dataset": "inventory"})["rows"]


@pytest.mark.django_db
def test_deliberate_retention_cutoff_and_permanent_reservation(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b)
    with clock.at(T + 100):
        run(
            b,
            "sale",
            {"animal": "A" * 16, "destination_pic": "DST", "nvd": "NVD", "price": 200},
        )
        with pytest.raises(ValidationError):
            run(a, "purge", {"animal": "A" * 16, "reason": "Too soon", "confirm": True})
    with clock.at(clock.calendar_days(T + 100, 366 * 11, "Australia/Melbourne")):
        with pytest.raises(PermissionDenied):
            run(b, "purge", {"animal": "A" * 16, "reason": "Test", "confirm": True})
        run(
            a,
            "purge",
            {
                "animal": "A" * 16,
                "reason": "Eligible detailed records removed",
                "confirm": True,
            },
        )
        assert Animal.objects.get().purged
        with pytest.raises(ValidationError):
            arrival(b)
        assert Identifier.objects.filter(value="A" * 16).exists()


@pytest.mark.django_db
def test_missing_purchase_cost_can_be_completed_without_rewriting_history(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b, purchase_cost="")
    with clock.at(T + 10):
        run(
            b,
            "edit_animal",
            {
                "animal": "A" * 16,
                "purchase_cost": "123.45",
                "breed": "Angus",
                "estimated_weeks": 1,
            },
        )
        row = reports.build(b, {})["rows"][0]
        assert row["financial_complete"] and row["purchase"] == D("123.45")
        assert Allocation.objects.filter(unknown=True).count() == 1


@pytest.mark.django_db
def test_historical_treatment_corrects_quantity_cost_and_restrictions(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 2)
        other = run(b, "mob", {"name": "Other"})["mob"]
        item = stock(b, kind="chemical", restriction_days=28)
    with clock.at(T + 100):
        run(b, "treat", {"mob": mob, "item": item, "dose": 2})
    with clock.at(T + 200):
        original_balance = Inventory.objects.get(pk=item).balance
        run(
            a,
            "move",
            {
                "animal": "EDGE000000000000",
                "destination": other,
                "effective_at": T + 50,
                "administration": True,
            },
        )
        assert Inventory.objects.get(pk=item).balance == original_balance + 2
        removed = s.find_animal(b, "EDGE000000000000")
        assert not removed.restrictions.filter(superseded_by__isnull=True).exists()
        assert (
            sum(removed.allocations.filter(category="chemical").values_list("amount", flat=True))
            == 0
        )
        kept = s.find_animal(b, "EDGE000000000001")
        assert (
            sum(kept.allocations.filter(category="chemical").values_list("amount", flat=True)) == -1
        )


@pytest.mark.django_db
def test_receipt_reversal_refuses_later_consumption_conflict(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 1)
        receipt = run(
            b,
            "receipt",
            {
                "kind": "feed",
                "name": "Only receipt",
                "packages": 1,
                "package_size": 330,
                "package_price": 65,
                "stock_unit": "kg",
            },
        )
    with clock.at(T + 1):
        run(
            b,
            "feed",
            {"mob": mob, "item": receipt["item"], "quantity": "0.5", "mode": "bales"},
        )
        with pytest.raises(ValidationError):
            run(
                a,
                "reverse",
                {
                    "original": receipt["event"],
                    "reason": "Cannot remove consumed stock",
                },
            )
        assert Inventory.objects.get(pk=receipt["item"]).balance == 165


@pytest.mark.django_db
def test_admin_identifier_review_never_merges_or_releases_reservations(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b, tag="A" * 16, rfid="000000000000001")
        arrival(b, tag="B" * 16, rfid="000000000000002")
        data = {
            "animal": "A" * 16,
            "identifier_kind": "rfid",
            "identifier_value": "000000000000002",
            "resolution": "reject_conflicting_link",
            "reason": "Checked source paperwork; the scanned tag belongs to the other animal",
        }
        with pytest.raises(PermissionDenied):
            run(b, "identifier_review", data)
        result = run(a, "identifier_review", data)
        assert not result["reservation_transferred"] and Animal.objects.count() == 2
        assert s.find_animal(b, "000000000000002").label == "B" * 16


@pytest.mark.django_db
def test_repurchase_resets_retention_eligibility(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b)
    with clock.at(T + 100):
        run(
            b,
            "sale",
            {"animal": "A" * 16, "destination_pic": "DST", "nvd": "NVD", "price": 200},
        )
    later = clock.calendar_days(T, 366 * 11, "Australia/Melbourne")
    with clock.at(later):
        arrival(b, entry_source="repurchase", purchase_cost=50)
        with pytest.raises(ValidationError):
            run(
                a,
                "purge",
                {"animal": "A" * 16, "reason": "Repurchased", "confirm": True},
            )
    with clock.at(later + 100):
        run(
            b,
            "sale",
            {
                "animal": "A" * 16,
                "destination_pic": "DST",
                "nvd": "NEW NVD",
                "price": 250,
            },
        )
        with pytest.raises(ValidationError):
            run(
                a,
                "purge",
                {
                    "animal": "A" * 16,
                    "reason": "Latest exit is recent",
                    "confirm": True,
                },
            )


@pytest.mark.django_db
def test_exited_attribute_and_tag_locks_cannot_be_bypassed_with_backdating(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b)
    with clock.at(T + 10):
        run(b, "sale", {"animal": "A" * 16, "destination_pic": "DST", "nvd": "NVD", "price": 200})
    with clock.at(T + 20):
        with pytest.raises(PermissionDenied):
            run(b, "edit_animal", {"animal": "A" * 16, "breed": "Changed", "effective_at": T + 1})
        with pytest.raises(PermissionDenied):
            run(
                b,
                "replace_tag",
                {"animal": "A" * 16, "nlis": "N" * 16, "reason": "Locked", "effective_at": T + 1},
            )


@pytest.mark.django_db
def test_historical_purchase_correction_is_scoped_to_its_presence_cycle(farm):
    _, a, b, c = farm
    with clock.at(T):
        arrival(b, purchase_cost="")
    with clock.at(T + 100):
        run(b, "sale", {"animal": "A" * 16, "destination_pic": "DST", "nvd": "FIRST", "price": 200})
    with clock.at(T + 200):
        arrival(b, entry_source="repurchase", purchase_cost=150)
    with clock.at(T + 300):
        run(
            a,
            "edit_animal",
            {
                "animal": "A" * 16,
                "effective_at": T + 50,
                "purchase_cost": "100",
                "reason": "Resolve the first purchase, retaining repurchase cost",
            },
        )
        row = reports.build(b, {})["rows"][0]
        assert row["purchase"] == 250 and row["revenue"] == 200 and row["financial_complete"]


@pytest.mark.django_db
def test_tied_second_membership_presence_and_repurchase_order(farm):
    _, a, b, c = farm
    with clock.at(T):
        mob = setup_mob(b, 1)
        item = stock(b)
        other = run(b, "mob", {"name": "Second-order mob"})["mob"]
        feeding = run(b, "feed", {"mob": mob, "item": item, "quantity": 1, "mode": "bales"})
        feed_event = Event.objects.get(pk=feeding["event"])
        animal = s.find_animal(b, "EDGE000000000000")
        run(b, "move", {"animal": animal.label, "destination": other})
        assert [x.id for x in s.members(b, mob, T, feed_event.effective_order)] == [animal.id]
        run(
            b,
            "sale",
            {"animal": animal.label, "destination_pic": "DST", "nvd": "NVD", "price": 200},
        )
        assert s.present(animal, T, feed_event.effective_order) and not s.present(animal, T)
        arrival(b, tag=animal.label, entry_source="repurchase", purchase_cost=50)
        assert s.present(animal, T)
